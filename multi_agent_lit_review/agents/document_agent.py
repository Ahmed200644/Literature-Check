import os
import datetime
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from tools.document_tools import set_cell_background, add_heading_with_spacing

class DocumentAgent:
    """
    DocumentAgent: Specialist agent responsible for compiling validated literature
    and synthesis into publication-grade Word (.docx) reports.
    """
    def __init__(self, output_dir: str = "."):
        self.output_dir = output_dir

    def create_literature_collection(self, papers: List[Dict[str, Any]], filename: str = "literature_collection_report.docx") -> str:
        """
        Generates the Literature Collection Report Word Document.
        """
        output_path = os.path.join(self.output_dir, filename)
        doc = Document()

        # Document Header
        h0 = doc.add_heading('Literature Collection Report (Week 1–5 Literature Agent)', level=0)
        h0.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p_meta = doc.add_paragraph()
        p_meta.add_run(f"Generated Date: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        p_meta.add_run(f"Total Validated Papers: {len(papers)}\n")
        p_meta.add_run("Target Problem: Agentic AI for Mathematical Discovery & ODE Predator–Prey Dynamics")

        add_heading_with_spacing(doc, "1. Annotated Literature Collection", level=1)

        table = doc.add_table(rows=1, cols=5)
        hdr_cells = table.rows[0].cells
        headers = ["Title", "Authors & Year", "DOI / Source", "Category", "Relevance & Key Summary"]
        widths = [Inches(2.2), Inches(1.5), Inches(1.5), Inches(1.2), Inches(2.6)]

        for i, (h, w) in enumerate(zip(headers, widths)):
            hdr_cells[i].text = h
            hdr_cells[i].width = w
            set_cell_background(hdr_cells[i], "2B5B84")
            for run in hdr_cells[i].paragraphs[0].runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

        for paper in papers:
            row_cells = table.add_row().cells
            for cell, w in zip(row_cells, widths):
                cell.width = w

            row_cells[0].text = paper.get("title", "Untitled")
            row_cells[1].text = f"{paper.get('authors', 'N/A')}\n({paper.get('year', 'N/A')})"
            row_cells[2].text = f"{paper.get('doi', 'N/A')}\n[{paper.get('data_source', 'Literature State')}]"
            row_cells[3].text = paper.get("category", "General")
            
            summary_text = paper.get("summary") or paper.get("abstract") or paper.get("relevance_reason") or ""
            row_cells[4].text = summary_text[:250] + ("..." if len(summary_text) > 250 else "")

        doc.save(output_path)
        print(f"[DocumentAgent] Saved Literature Collection Report to: {output_path}")
        return output_path

    def create_review_paper_draft(self, papers: List[Dict[str, Any]], synthesis: Dict[str, Any], filename: str = "review_paper_draft.docx") -> str:
        """
        Generates the Structured Review Paper Draft Word Document.
        """
        output_path = os.path.join(self.output_dir, filename)
        doc = Document()

        # Title
        title_p = doc.add_heading('Agentic AI Systems for Autonomous Mathematical Discovery & ODE Ecological Modeling: A Literature Review', level=0)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_p.add_run("Egypt Scholar Advanced Lab 12 — Team 7\n").bold = True
        meta_p.add_run("Supervisor: Dr. Omar A. M. Abdelraouf (Ain Shams University)\n")
        meta_p.add_run(f"Date: {datetime.datetime.now().strftime('%B %Y')}\n\n")

        # Abstract / Executive Summary
        add_heading_with_spacing(doc, "Abstract", level=1)
        doc.add_paragraph(
            "This review synthesizes modern literature at the intersection of Agentic AI, Autonomous Mathematical "
            "Discovery, and Ordinary Differential Equation (ODE) formulation, with an emphasis on dynamical and ecological "
            "systems such as Predator–Prey and Lotka–Volterra population dynamics. Based on a validated literature corpus "
            f"of {len(papers)} papers, we classify existing frameworks into direct LLM modeling agents, methodological "
            "dynamical systems tools, and adjacent autonomous control architectures."
        )

        # Section 1: Introduction
        add_heading_with_spacing(doc, "1. Introduction & Research Scope", level=1)
        doc.add_paragraph(
            "Mathematical formulation of complex non-linear dynamical systems historically relies on expert intuition and manual "
            "parameter estimation. With the emergence of Large Language Model (LLM) agents, automated scientific discovery "
            "has evolved from simple prompt engineering to context engineering, harness engineering, and autonomous AI loops."
        )

        # Section 2: Thematic Analysis
        add_heading_with_spacing(doc, "2. Thematic Classification of Validated Literature", level=1)

        categories = synthesis.get("categories", {})
        for cat_name, cat_papers in categories.items():
            add_heading_with_spacing(doc, f"2.{list(categories.keys()).index(cat_name) + 1} {cat_name} Research Frameworks (N={len(cat_papers)})", level=2)
            for p in cat_papers[:10]:
                p_item = doc.add_paragraph(style='List Bullet')
                r_title = p_item.add_run(f"{p.get('title')} ")
                r_title.bold = True
                p_item.add_run(f"({p.get('authors')}, {p.get('year')}). ")
                p_item.add_run(p.get('summary', p.get('abstract', ''))[:200] + "...")

        # Section 3: Key Research Gaps & Future Directions
        add_heading_with_spacing(doc, "3. Research Gaps & Future Directions", level=1)
        doc.add_paragraph(synthesis.get("gap_analysis", "Autonomous ODE discovery under sparse observational data remains an open challenge."))

        doc.save(output_path)
        print(f"[DocumentAgent] Saved Review Paper Draft to: {output_path}")
        return output_path
