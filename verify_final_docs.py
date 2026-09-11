import os
import sys
import docx

def find_docx_file(filename="review_paper_draft.docx"):
    # Check current directory
    if os.path.exists(filename):
        return filename
    
    # Check multi_agent_lit_review directory relative to current working directory
    multi_agent_path = os.path.join("multi_agent_lit_review", filename)
    if os.path.exists(multi_agent_path):
        return multi_agent_path

    # Check relative to script directory
    script_dir = os.path.dirname(os.path.abspath(__file__))
    script_relative = os.path.join(script_dir, filename)
    if os.path.exists(script_relative):
        return script_relative

    script_multi_agent = os.path.join(script_dir, "multi_agent_lit_review", filename)
    if os.path.exists(script_multi_agent):
        return script_multi_agent

    raise FileNotFoundError(f"Could not locate '{filename}'")

def main():
    docx_path = find_docx_file("review_paper_draft.docx")
    print(f"Opening docx file: {docx_path}")
    
    doc = docx.Document(docx_path)
    
    # Count paragraphs and headings
    paragraph_count = len(doc.paragraphs)
    heading_count = sum(1 for p in doc.paragraphs if p.style and p.style.name.startswith("Heading"))
    
    print(f"Number of paragraphs: {paragraph_count}")
    print(f"Number of heading paragraphs: {heading_count}")
    return paragraph_count

if __name__ == "__main__":
    main()
